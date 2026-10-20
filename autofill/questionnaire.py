"""Read questionnaires (Google VSAQ JSON or Excel) and write the filled-in Excel."""
import html
import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

QUESTION_TYPES = {"yesno", "line", "box", "radiogroup", "checkgroup"}
QUESTION_HEADERS = ("question", "questions", "control question", "query")


def _strip_html(s: str) -> str:
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", s))).strip()


def load_vsaq(path: Path) -> list[dict]:
    """Flatten a VSAQ questionnaire into [{id, section, question}]. VSAQ files start with // comments."""
    raw = "\n".join(l for l in path.read_text(encoding="utf-8").splitlines() if not l.lstrip().startswith("//"))
    raw = re.sub(r'"\s*\+\s*"', "", raw)  # VSAQ is JS-flavoured: joins long strings with +
    data = json.loads(raw)
    out = []

    def walk(node, section):
        if isinstance(node, list):
            for n in node:
                walk(n, section)
        elif isinstance(node, dict):
            if node.get("type") == "block" and node.get("text"):
                section = _strip_html(node["text"])
            if node.get("type") in QUESTION_TYPES and node.get("id") and node.get("text"):
                q = _strip_html(node["text"])
                if q.endswith("?") or node["type"] == "yesno":  # skip free-text labels that are not questions
                    out.append({"id": node["id"], "section": section, "question": q})
            for v in node.values():
                walk(v, section)

    walk(data, "")
    return out


def read_xlsx(path) -> list[dict]:
    """Find the question column by header name and return [{id, section, question}]."""
    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    for h, header in enumerate(rows[:10]):
        cells = [str(c).strip().lower() if c else "" for c in header]
        qcol = next((i for i, c in enumerate(cells) if c in QUESTION_HEADERS), None)
        if qcol is not None:
            idcol = next((i for i, c in enumerate(cells) if c in ("id", "#", "ref", "question id")), None)
            seccol = next((i for i, c in enumerate(cells) if c in ("section", "domain", "category")), None)
            return [
                {
                    "id": str(r[idcol]) if idcol is not None and r[idcol] else str(n + 1),
                    "section": str(r[seccol]) if seccol is not None and r[seccol] else "",
                    "question": str(r[qcol]).strip(),
                }
                for n, r in enumerate(rows[h + 1:])
                if r[qcol] and str(r[qcol]).strip()
            ]
    raise ValueError("No column headed 'Question' found in the first 10 rows")


def write_questions_xlsx(questions: list[dict], path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questionnaire"
    ws.append(["ID", "Section", "Question"])
    for q in questions:
        ws.append([q["id"], q["section"], q["question"]])
    _style(ws, [14, 30, 90])
    wb.save(path)


def write_answers_xlsx(results: list[dict], path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Answers"
    ws.append(["ID", "Question", "Answer", "Explanation", "Confidence", "Status", "Citations"])
    review = PatternFill("solid", fgColor="FFF4CE")
    for r in results:
        cites = "\n".join(f"{c['title']} > {c['section']} ({c['url']})" for c in r["citations"])
        ws.append([r["id"], r["question"], r["answer"], r["explanation"], r["confidence"], r["status"], cites])
        if r["status"] != "Drafted":
            for cell in ws[ws.max_row]:
                cell.fill = review
    _style(ws, [12, 50, 12, 70, 12, 16, 70])
    wb.save(path)


def _style(ws, widths):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
